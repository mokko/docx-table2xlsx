"""
Quick and dirty converter that turns docx tables into simple Excel sheets

Will mercylessly overwrite output file 

For MDS

Commandline Usage
    t2x.py -i infile.docx
"""
import sys
import os
import xlsxwriter
import inspect
from docx.api import Document

class Table2xlsx:
    #from docx.api import Document
    #from docx import Document

    def __init__(self, in_fn):
        out_fn = os.path.splitext(in_fn)[0] + ".xlsx"

        document = Document(in_fn)
        table = document.tables[0]
        wb = xlsxwriter.Workbook(out_fn) #check if file is open here?
        ws = wb.add_worksheet("t2x")

        #for shape in document.inline_shapes:
        #    print(shape)

        r=0
        for row in table.rows:
            print (f"{r}\b\b\b\b", end='', flush=True) #very quick and very dirty
            c=0
            for cell in row.cells:
                if not cell.text:
                    print (f"xxx {r}:{c}")
                ws.write(r,c, cell.text)
                c=c+1
            r=r+1
        wb.close()
    
if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('-i', '--input', required=True)
    args = parser.parse_args()
    t=Table2xlsx (args.input)
